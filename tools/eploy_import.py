"""Eploy import: turns RAC's application report into the aggregated counts the
planner uses for screening pass rates and hire rates.

Usage (from the repo folder):
    python tools/eploy_import.py "PATH/TO/Eploy workbook.xlsx"
    python tools/eploy_import.py "PATH/TO/Eploy workbook.xlsx" --check
    python tools/eploy_import.py "PATH/TO/Eploy workbook.xlsx" --accept-changes

What it does, in order. It stops with a message at the first problem:
1. Reads the sheet "Application Report - All Roles" by column header name
   (row 2), not by position. A missing required column stops the import.
2. Maps roles, regions and sources using data/eploy_mappings.csv. Any label not
   in that file stops the import and is listed, so a new label is decided on,
   not guessed.
3. Validates every planned-role row: Hired and Progressed Past Screening must
   read TRUE or FALSE (an empty value usually means formulas without saved
   values: open and save the file in Excel first); application dates must fall
   between eploy_first_month and the file's own date.
4. Compares with the previous import (data/eploy_rates.json) on the months both
   hold: applications, progressed and hires by role and platform. Any shift
   above 10% on a count of 20 or more stops the import unless
   --accept-changes is given.
5. Writes data/eploy_rates.json: counts by role, region, platform and
   application month, plus the dataset file name and date. No candidate-level
   field is written.

--check runs steps 1 to 4 and confirms data/eploy_rates.json already matches
the workbook, without writing anything.

The workbook itself holds candidate-level data and is never committed.
"""
import argparse, csv, datetime, hashlib, json, os, sys
from collections import defaultdict

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
MAPPINGS = os.path.join(ROOT, 'data', 'eploy_mappings.csv')
OUT = os.path.join(ROOT, 'data', 'eploy_rates.json')
ASSUMPTIONS = os.path.join(ROOT, 'assumptions.csv')
SHEET = 'Application Report - All Roles'
HEADER_ROW = 2
REQUIRED = {
    'date': 'Application Date',
    'source': 'Corrected Detail',
    'hired': 'Hired',
    'progressed': 'Progressed Past Screening',
    'role': 'Patrol or SMR',
    'region': 'Region',
}
PLATFORMS = ('indeed', 'meta', 'google', 'appcast', 'other')
SHIFT_LIMIT = 0.10
SHIFT_MIN_COUNT = 20


class ImportStopped(Exception):
    pass


def norm(label):
    """Labels are matched trimmed and case-insensitive; blank is '(blank)'."""
    s = '' if label is None else str(label).strip()
    return '(blank)' if s == '' else s.lower()


def read_mappings(path=MAPPINGS):
    maps = {'role': {}, 'region': {}, 'source': {}}
    with open(path, newline='', encoding='utf-8') as f:
        for row in csv.DictReader(f):
            kind = row['type'].strip()
            if kind not in maps:
                raise ImportStopped(f'{os.path.basename(path)}: unknown mapping type "{kind}"')
            key = norm(row['label'])
            if key in maps[kind]:
                raise ImportStopped(f'{os.path.basename(path)}: {kind} label "{row["label"]}" is listed twice')
            target = row['maps_to'].strip()
            if kind == 'source' and target not in PLATFORMS:
                raise ImportStopped(f'{os.path.basename(path)}: source "{row["label"]}" maps to "{target}", expected one of {", ".join(PLATFORMS)}')
            maps[kind][key] = target
    return maps


def first_month():
    with open(ASSUMPTIONS, newline='', encoding='utf-8') as f:
        for row in csv.DictReader(f):
            if row['key'] == 'eploy_first_month':
                return row['value'].strip()
    raise ImportStopped('assumptions.csv has no eploy_first_month row')


def as_flag(value, column, problems):
    if value is True or (isinstance(value, str) and value.strip().upper() == 'TRUE'):
        return 1
    if value is False or (isinstance(value, str) and value.strip().upper() == 'FALSE'):
        return 0
    key = 'empty' if value is None or str(value).strip() == '' else repr(value)
    problems[(column, key)] += 1
    return None


def read_workbook(path, maps, start_month, file_date):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        return read_sheet(wb, maps, start_month, file_date)
    finally:
        wb.close()   # read-only workbooks hold the file open until closed


def read_sheet(wb, maps, start_month, file_date):
    if SHEET not in wb.sheetnames:
        raise ImportStopped(f'sheet "{SHEET}" not found; sheets were: {", ".join(wb.sheetnames)}')
    rows = wb[SHEET].iter_rows(values_only=True)
    header = None
    for i, r in enumerate(rows, start=1):
        if i == HEADER_ROW:
            header = [('' if h is None else str(h).strip()) for h in r]
            break
    if header is None:
        raise ImportStopped('the sheet has no header row')
    col = {}
    missing = []
    for key, name in REQUIRED.items():
        hits = [i for i, h in enumerate(header) if h == name]
        if not hits:
            missing.append(name)
        elif len(hits) > 1:
            raise ImportStopped(f'column "{name}" appears {len(hits)} times in the header row')
        else:
            col[key] = hits[0]
    if missing:
        raise ImportStopped('required columns not found in row 2: ' + ', '.join(missing))

    unmapped = defaultdict(int)
    problems = defaultdict(int)
    counts = defaultdict(lambda: [0, 0, 0])
    stats = defaultdict(int)
    first = last = None
    lo = start_month + '-01'
    hi = file_date
    for r in rows:
        if r is None or all(v is None for v in r):
            continue
        stats['rows_read'] += 1
        role_label = r[col['role']]
        if norm(role_label) not in maps['role']:
            unmapped[('role', '(blank)' if norm(role_label) == '(blank)' else str(role_label).strip())] += 1
            continue
        role = maps['role'][norm(role_label)]
        if not role:
            stats['rows_other_roles'] += 1
            continue
        stats['rows_' + role] += 1
        region_label = r[col['region']]
        source_label = r[col['source']]
        region = maps['region'].get(norm(region_label))
        platform = maps['source'].get(norm(source_label))
        if region is None:
            unmapped[('region', str(region_label).strip())] += 1
        if platform is None:
            unmapped[('source', str(source_label).strip())] += 1
        d = r[col['date']]
        if isinstance(d, datetime.datetime):
            day = d.date().isoformat()
        elif isinstance(d, datetime.date):
            day = d.isoformat()
        else:
            problems[('Application Date', 'not a date: ' + repr(d)[:30])] += 1
            continue
        if not (lo <= day <= hi):
            problems[('Application Date', f'outside {lo} to {hi}')] += 1
            continue
        hired = as_flag(r[col['hired']], REQUIRED['hired'], problems)
        prog = as_flag(r[col['progressed']], REQUIRED['progressed'], problems)
        if region is None or platform is None or hired is None or prog is None:
            continue
        if hired and not prog:
            stats['hired_not_marked_progressed_' + role] += 1
        first = day if first is None or day < first else first
        last = day if last is None or day > last else last
        c = counts[(role, region, platform, day[:7])]
        c[0] += 1
        c[1] += prog
        c[2] += hired
    if unmapped:
        lines = [f'  {kind}: "{label}" ({n} rows)' for (kind, label), n in sorted(unmapped.items())]
        raise ImportStopped('labels not in data/eploy_mappings.csv; add each one, then re-run:\n' + '\n'.join(lines))
    if problems:
        lines = [f'  {column}: {what} ({n} rows)' for (column, what), n in sorted(problems.items())]
        hint = ''
        if any(what == 'empty' for (_, what) in problems):
            hint = '\nEmpty values usually mean formulas without saved values: open the workbook in Excel, save it, and re-run.'
        raise ImportStopped('unexpected values:\n' + '\n'.join(lines) + hint)
    return counts, dict(stats), first, last


def totals(cells, months=None):
    t = defaultdict(lambda: [0, 0, 0])
    for role, region, platform, month, a, p, h in cells:
        if months is not None and month not in months:
            continue
        for key in ((role, platform), (role, 'all sources')):
            t[key][0] += a
            t[key][1] += p
            t[key][2] += h
    return t


def compare(previous, cells):
    """Shifts above the limit on the months both imports hold."""
    old_months = {c[3] for c in previous['cells']}
    new_months = {c[3] for c in cells}
    common = old_months & new_months
    old, new = totals(previous['cells'], common), totals(cells, common)
    shifts, notes = [], []
    for key in sorted(set(old) | set(new)):
        for i, what in enumerate(('applications', 'progressed', 'hires')):
            a, b = old[key][i], new[key][i]
            if a == b:
                continue
            change = (b - a) / a if a else float('inf')
            line = f'  {key[0]} {key[1]} {what}: {a} to {b} ({change:+.1%})' if a else f'  {key[0]} {key[1]} {what}: {a} to {b}'
            if abs(change) > SHIFT_LIMIT and max(a, b) >= SHIFT_MIN_COUNT:
                shifts.append(line)
            else:
                notes.append(line)
    return shifts, notes, sorted(new_months - old_months), sorted(old_months - new_months)


def build(path, accept_changes=False):
    maps = read_mappings()
    start = first_month()
    file_date = datetime.datetime.fromtimestamp(os.path.getmtime(path)).date().isoformat()
    counts, stats, first, last = read_workbook(path, maps, start, file_date)
    cells = [[*k, *v] for k, v in sorted(counts.items())]
    with open(MAPPINGS, 'rb') as f:   # line endings ignored, so a Windows checkout matches
        mapping_hash = hashlib.sha256(f.read().replace(b'
', b'
')).hexdigest()[:12]
    unknown = defaultdict(int)
    for role, region, platform, month, a, p, h in cells:
        if region == 'Unknown':
            unknown[role] += a
    out = {
        'note': 'Aggregated from the Eploy application report by tools/eploy_import.py. '
                'cells: [role, region, platform, application month, applications, progressed past screening, hires].',
        'dataset': {
            'file': os.path.basename(path),
            'file_date': file_date,
            'sheet': SHEET,
            'first_application': first,
            'last_application': last,
            'rows_read': stats.get('rows_read', 0),
            'rows_smr': stats.get('rows_SMR', 0),
            'rows_patrol': stats.get('rows_Patrol', 0),
            'rows_other_roles': stats.get('rows_other_roles', 0),
            'hired_not_marked_progressed': {r: stats.get('hired_not_marked_progressed_' + r, 0) for r in ('SMR', 'Patrol')},
            'unknown_region_applications': dict(unknown),
        },
        'mappings_sha256': mapping_hash,
        'eploy_first_month': start,
        'cells': cells,
    }
    report = [f'Read {out["dataset"]["rows_read"]} rows: SMR {out["dataset"]["rows_smr"]}, Patrol {out["dataset"]["rows_patrol"]}, '
              f'other roles {out["dataset"]["rows_other_roles"]} (not imported).',
              f'Applications {first} to {last}; file dated {file_date}.',
              f'Applications with no region (role totals only): {dict(unknown)}.',
              f'Hired but not marked progressed: {out["dataset"]["hired_not_marked_progressed"]}.']
    previous = None
    if os.path.exists(OUT):
        with open(OUT, encoding='utf-8') as f:
            previous = json.load(f)
    if previous:
        shifts, notes, added, dropped = compare(previous, cells)
        report.append(f'Compared with the previous import ({previous["dataset"]["file"]}, {previous["dataset"]["file_date"]}): '
                      f'months added {added or "none"}, months no longer held {dropped or "none"}.')
        if notes:
            report.append('Smaller changes on shared months:\n' + '\n'.join(notes))
        if shifts:
            msg = 'Shifts above 10% on months both imports hold:\n' + '\n'.join(shifts)
            if not accept_changes:
                raise ImportStopped(msg + '\nCheck the new file. If the changes are expected, re-run with --accept-changes.')
            report.append(msg + '\nAccepted (--accept-changes).')
    else:
        report.append('No previous import to compare with.')
    return out, previous, report


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('workbook')
    ap.add_argument('--check', action='store_true', help='validate and confirm data/eploy_rates.json matches; write nothing')
    ap.add_argument('--accept-changes', action='store_true', help='write even if counts shifted more than 10%%')
    args = ap.parse_args(argv)
    try:
        out, previous, report = build(args.workbook, args.accept_changes)
    except ImportStopped as e:
        print('IMPORT STOPPED: ' + str(e))
        return 1
    print('\n'.join(report))
    if args.check:
        same = previous is not None and previous['cells'] == out['cells'] and previous['dataset'] == out['dataset'] \
            and previous['mappings_sha256'] == out['mappings_sha256']
        print('CHECK: data/eploy_rates.json matches the workbook' if same
              else 'CHECK FAILED: data/eploy_rates.json does not match the workbook; run the import without --check')
        return 0 if same else 1
    with open(OUT, 'w', encoding='utf-8', newline='\n') as f:
        f.write('{\n')
        items = [(k, v) for k, v in out.items() if k != 'cells']
        for k, v in items:
            f.write(f'  {json.dumps(k)}: {json.dumps(v)},\n')
        f.write('  "cells": [\n')
        f.write(',\n'.join('    ' + json.dumps(c) for c in out['cells']))
        f.write('\n  ]\n}\n')
    print(f'Wrote {os.path.relpath(OUT, ROOT)}: {len(out["cells"])} aggregated rows.')
    return 0


if __name__ == '__main__':
    sys.exit(main())
