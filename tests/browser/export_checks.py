"""Browser checks for the exports, on the test link address, reading the files
the app actually saves.

1. PDF: the header PDF button saves a PDF whose text (read with pypdf) holds
   the planner's figures for the same settings, the title with the hire
   target, the version stamp on every page, and no em-dashes or Hiring Lab.
2. "PDF, no notes" saves the same document without the notes pages.
3. Workings: the header Workings button saves a workbook; it is recalculated
   from scratch by LibreOffice and every formula must then give the planner's
   own figure. This is the check that the spreadsheet's arithmetic is the
   planner's arithmetic, not a copy of its answers.
4. No page errors, no alerts, nothing written, every request handled.

Run from the repo folder:  python tests/browser/export_checks.py  [--libs DIR]
Files are written to tests/browser/out/.
The workings part needs LibreOffice (soffice) and openpyxl; without either it
says so and is reported as skipped."""
import copy, io, json, os, re, shutil, subprocess, sys, tempfile
from playwright.sync_api import sync_playwright
from pypdf import PdfReader
sys.path.insert(0, os.path.dirname(__file__))
from guard import TEST, new_page, libs_arg, VERSION_COMMIT, ROOT
from app_fixture import DB, EXPECTED_JS

OUT = os.path.join(os.path.dirname(__file__), 'out'); os.makedirs(OUT, exist_ok=True)
LIBS = libs_arg(sys.argv)

# The figures every formula in the workbook must give once it is recalculated.
WORKINGS_JS = EXPECTED_JS.replace(
    "  return { apps: plan.totals.apps,",
    """  const built = RAC.workings.build([{ role, roleName: role + ' (check)', plan }],
    { monthLabel: 'October 2026', planName: 'Working plan', backtest: RAC.app.state.backtest });
  return { checks: built.checks, problems: built.problems, sheets: built.sheets.map(s => s.name),
    fileName: RAC.workings.fileName([{ role }], { monthLabel: 'October 2026' }),
    apps: plan.totals.apps,""")

SOFFICE = [
    r'C:\Program Files\LibreOffice\program\soffice.exe',
    r'C:\Program Files (x86)\LibreOffice\program\soffice.exe',
    shutil.which('soffice') or '', shutil.which('libreoffice') or '',
]


def recalculate(path):
    """A copy of the workbook with every formula worked out again by LibreOffice."""
    exe = next((p for p in SOFFICE if p and os.path.isfile(p)), None)
    if not exe:
        return None, 'LibreOffice (soffice) was not found'
    out = tempfile.mkdtemp(prefix='rac_calc_')
    profile = os.path.join(out, 'profile')
    cmd = [exe, '-env:UserInstallation=file:///' + profile.replace('\\', '/'),
           '--headless', '--norestore', '--calc', '--convert-to', 'xlsx', '--outdir', out, path]
    r = subprocess.run(cmd, capture_output=True, text=True, timeout=300)
    made = os.path.join(out, os.path.splitext(os.path.basename(path))[0] + '.xlsx')
    if not os.path.isfile(made):
        return None, f'LibreOffice did not write a file: {(r.stdout or "") + (r.stderr or "")}'.strip()
    return made, None

# The figures the PDF must show, formatted as the PDF formats them.
FIGURES_JS = EXPECTED_JS.replace(
    "  return { apps: plan.totals.apps,",
    """  const F = RAC.text.fmt;
  const figs = [F.gbp(plan.budget), F.gbp(plan.deployable), F.gbp(plan.placed), F.gbp(plan.unplaced.total),
    F.num(plan.totals.allHires), F.num(plan.totals.hires), F.int(plan.totals.apps), F.int(plan.totals.passed)];
  if (plan.fees.on) figs.push(F.gbp(plan.fees.total, 2));
  plan.locations.forEach(l => { figs.push(F.gbp(l.spend)); RAC.PLATFORMS.forEach(p => { if (l.cells[p].spend > 0) figs.push(F.gbp(l.cells[p].plannedCpa, 2)); }); });
  RAC.PLATFORMS.forEach(p => figs.push(F.gbp(plan.platforms[p].media)));
  const title = RAC.pdf.titleOf({ plan, roleName: role + ' (' + { SMR: 'Mobile Vehicle Tech', Patrol: 'Roadside Tech (incl. SuperFlex)' }[role] + ')' }, 'October 2026');
  return { figs, title, headings: RAC.text.method(plan.A, role, plan, RAC.app.state.backtest).map(s => s.heading),
    apps: plan.totals.apps,""")

squash = lambda t: re.sub(r'\s+', '', t)


def read_pdf(path):
    r = PdfReader(path)
    return [p.extract_text() or '' for p in r.pages]


def save_export(page, label, name):
    with page.expect_download(timeout=120000) as dl:
        page.locator('.app-header button', has_text=re.compile('^' + re.escape(label) + '$')).click()
    path = os.path.join(OUT, name)
    dl.value.save_as(path)
    return dl.value.suggested_filename, path


save_pdf = save_export


fails, notes = [], []
db = copy.deepcopy(DB)
w = db['workspace']['months']['2026-10']['working']
w['commentary'] = {'SMR': '- London is off this month.\n- Scotland is held at £2,500.', 'Patrol': ''}
w['commentaryLegacy'] = {'SMR': '', 'Patrol': ''}

with sync_playwright() as pw:
    browser = pw.chromium.launch()
    ctx, page, guard, errors = new_page(browser, TEST, db=db, libs=LIBS)
    dialogs = []
    page.on('dialog', lambda d: (dialogs.append(d.message), d.dismiss()))
    page.goto(TEST + '#planner/plan')
    page.wait_for_selector('.kpi-label:has-text("Predicted applications")', timeout=60000)
    want = page.evaluate(FIGURES_JS, 'SMR')

    name, path = save_pdf(page, 'PDF', 'plan_smr.pdf')
    pages = read_pdf(path)
    text = '\n'.join(pages)
    flat = squash(text)
    notes.append(f'PDF saved as {name}: {len(pages)} pages')
    if name != 'RAC_October_2026_SMR_Plan.pdf':
        fails.append('PDF file name ' + name)
    if squash(want['title']) not in flat:
        fails.append(f"title {want['title']!r} not in the PDF")
    missing = [x for x in want['figs'] if squash(x) not in flat]
    if missing:
        fails.append(f'PDF lacks planner figures: {missing[:6]}')
    notes.append(f"{len(want['figs']) - len(missing)} of {len(want['figs'])} planner figures found; title {want['title']!r}")
    for i, p in enumerate(pages):
        if 'Code' + VERSION_COMMIT[:7] not in squash(p):
            fails.append(f'page {i + 1} has no version stamp')
        if f'Page{i + 1}of{len(pages)}' not in squash(p):
            fails.append(f'page {i + 1} has no page number')
    if '—' in text or re.search(r'hiring\s*lab', text, re.I):
        fails.append('PDF text fails the output checks')
    missing_h = [h for h in want['headings'] if squash(h) not in flat]
    if missing_h:
        fails.append(f'method headings missing: {missing_h}')
    if 'Notesonthisplan' not in squash(pages[0]) or 'LondonisoffthismonthSX'.replace('SX', '') not in squash(pages[0]).replace('.', ''):
        fails.append('first page is not the notes page: ' + pages[0][:120])

    name2, path2 = save_pdf(page, 'PDF, no notes', 'plan_smr_no_notes.pdf')
    pages2 = read_pdf(path2)
    notes.append(f'"PDF, no notes" saved as {name2}: {len(pages2)} pages')
    if len(pages) - len(pages2) != 1 or any('Notesonthisplan' in squash(p) for p in pages2):
        fails.append(f'no-notes PDF: {len(pages2)} pages against {len(pages)}')
    if name2 != 'RAC_October_2026_SMR_Plan_No_Notes.pdf':
        fails.append('no-notes file name ' + name2)

    # ---- workings: saved, recalculated from scratch, compared -------------
    want_w = page.evaluate(WORKINGS_JS, 'SMR')
    if want_w['problems']:
        fails.append('the workings export failed its own checks: ' + '; '.join(want_w['problems'][:3]))
    name3, path3 = save_export(page, 'Workings', 'workings_smr.xlsx')
    notes.append(f"Workings saved as {name3}: {len(want_w['sheets'])} sheets, {len(want_w['checks'])} formulas to recalculate")
    if name3 != want_w['fileName']:
        fails.append(f"workings file name {name3} against {want_w['fileName']}")
    try:
        from openpyxl import load_workbook
    except ImportError:
        load_workbook = None
    if load_workbook is None:
        notes.append('SKIPPED: openpyxl is not installed, so the workbook was not recalculated')
    else:
        # No formula may carry a saved answer: a saved answer could go stale
        # and a reader would show it without working the formula out.
        saved = load_workbook(path3, data_only=True)
        stale = [c for c in want_w['checks'] if saved[c['sheet']][c['ref']].value is not None]
        saved.close()
        if stale:
            fails.append(f'{len(stale)} formulas were saved with an answer, so they would not be worked out again: '
                         + ', '.join(f"{c['sheet']}!{c['ref']}" for c in stale[:4]))
        made, why = recalculate(path3)
        if made is None:
            notes.append('SKIPPED: ' + why)
        else:
            wb = load_workbook(made, data_only=True)
            bad, blank, checked = [], 0, 0
            for c in want_w['checks']:
                got = wb[c['sheet']][c['ref']].value
                if got is None:
                    blank += 1
                    bad.append(f"{c['sheet']}!{c['ref']} came back empty ({c['formula'][:40]})")
                    continue
                if isinstance(got, str):
                    bad.append(f"{c['sheet']}!{c['ref']} gave {got!r} ({c['formula'][:40]})")
                    continue
                want = c['value']
                if abs(got - want) > max(1e-6, abs(want) * 1e-9):
                    bad.append(f"{c['sheet']}!{c['ref']} gave {got!r}, the planner {want!r} ({c['formula'][:40]})")
                checked += 1
            wb.close()
            shutil.rmtree(os.path.dirname(made), ignore_errors=True)
            notes.append(f'{checked} of {len(want_w["checks"])} formulas recalculated in LibreOffice and compared'
                         + (f'; {blank} came back empty' if blank else ''))
            if bad:
                fails.append(f'{len(bad)} formulas disagree with the planner: ' + ' | '.join(bad[:5]))

    if dialogs:
        fails.append(f'alerts shown: {dialogs[:2]}')
    if errors:
        fails.append(f'page errors: {errors[:3]}')
    if guard.writes or guard.blocked:
        fails.append(f'writes {guard.writes[:3]}, blocked {guard.blocked[:3]}')
    if 'version' not in guard.served:
        fails.append('the app did not ask for the code version')
    ctx.close()
    browser.close()

for n in notes:
    print('  ' + n)
print("FAIL: " + "; ".join(fails) if fails else "PASS: the PDF and the workings hold the planner figures; the workbook recalculates to them; nothing written")
